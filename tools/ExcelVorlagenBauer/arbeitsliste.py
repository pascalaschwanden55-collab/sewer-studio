# -*- coding: utf-8 -*-
"""Kompakte Bedienung bei unveraenderter Spalten- und Datenzeilenfolge."""
import stil as S
import werkzeug as W
from openpyxl.utils import get_column_letter as SP
from openpyxl.worksheet.properties import PageSetupProperties


def anlegen(wb, name, breiten, letzte):
    ws = wb.active
    ws.row_dimensions[24].height = 26
    ws.merge_cells("A24:B24")
    ws["A24"] = "Aktuelle Auswahl"
    ws["C24"] = "=SUBTOTAL(103,$A$27:$A$%d)" % letzte
    ws["C24"].number_format = '0" %s"' % ("Haltungen" if name == "Haltungen" else "Schächte")
    if name == "Haltungen":
        ws["D24"] = "Länge m"
        ws.merge_cells("E24:F24")
        ws["E24"] = "=SUBTOTAL(109,$G$27:$G$%d)" % letzte
        ws["E24"].number_format = S.FORMAT_METER
        kosten = "N"
    else:
        kosten = "I"
    ws["G24"] = "Kosten CHF"
    ws["H24"] = "=SUBTOTAL(109,$%s$27:$%s$%d)" % (kosten, kosten, letzte)
    ws["H24"].number_format = '#,##0.00'
    for row in ws.iter_rows(min_row=24, max_row=24, max_col=8):
        for cell in row:
            cell.font = S.schrift(10, fett=True)
            cell.fill = S.fuellung(S.BLOCK_FELD)
            cell.alignment = S.ausrichtung("left")
    for adresse in ["C24", "E24", "H24"]:
        ws[adresse].alignment = S.ausrichtung("right", einzug=1)
    return ws


def abschliessen(ws):
    kennung = 2 if ws.title == "Haltungen" else 3
    ws.freeze_panes = ws.cell(W.Z_DATEN, kennung + 1)
    ws.sheet_view.topLeftCell = "A1"
    # Alle Spalten gemeinsam auf einer Seitenbreite.
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.print_title_cols = "A:%s" % SP(kennung)
    ws.print_title_rows = "%d:%d" % (W.Z_TABKOPF, W.Z_TABKOPF)
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.page_margins.header = ws.page_margins.footer = 0.15
